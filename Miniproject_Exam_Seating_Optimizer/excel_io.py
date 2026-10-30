"""Excel import/export for the desktop app (openpyxl)."""
from pathlib import Path
import os
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from data import FIELDS, YEARS


def read_data(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    batches = {}
    try:
        for sheet in wb:
            table = sheet.title.strip().lower()
            if table not in FIELDS:
                continue
            if table in batches:
                raise ValueError(f"Duplicate data sheet: {table}")
            rows = sheet.iter_rows()
            header_row = next(rows, None)
            if header_row is None:
                raise ValueError(f"{table}: missing header row.")
            headers = [str(c.value or "").strip().lower() for c in header_row]
            required = FIELDS[table]
            if any(headers.count(key) != 1 for key in required):
                raise ValueError(f"{table}: row 1 must contain each header exactly once: {', '.join(required)}")
            indices = [headers.index(key) for key in required]
            data = []
            for row_index, cells in enumerate(rows, start=2):
                values = [cells[i].value if i < len(cells) else None for i in indices]
                if all(value is None or value == "" for value in values):
                    continue
                if any(i < len(cells) and cells[i].data_type == "f" for i in indices):
                    raise ValueError(f"{table}, Excel row {row_index}: use values, not formulas.")
                data.append(dict(zip(required, ["" if v is None else v for v in values])))
            batches[table] = data
        if not batches:
            raise ValueError("No data sheets found. Use students, rooms, subjects or supervisors.")
    finally:
        wb.close()
    return batches


def append_values(sheet, values):
    """Force user-supplied text to remain literal text in Excel."""
    sheet.append(values)
    for cell, value in zip(sheet[sheet.max_row], values):
        if isinstance(value, str):
            cell.data_type = "s"


def style_sheet(sheet):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="263D7A")
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="17233D")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F4FC")
        sheet.row_dimensions[row[0].row].height = 32
    for column in sheet.columns:
        width = max(len(str(c.value or "")) for c in column) + 3
        sheet.column_dimensions[column[0].column_letter].width = min(48, max(14, width))
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True


def save_workbook(wb, path):
    """Replace a completed export atomically, keeping an old file on failure."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=path.parent, delete=False) as temp:
            temp_path = Path(temp.name)
        wb.save(temp_path)
        os.replace(temp_path, path)
    finally:
        wb.close()
        if temp_path and temp_path.exists():
            temp_path.unlink()


def export_data(store, path):
    wb = Workbook()
    wb.remove(wb.active)
    for table, fields in FIELDS.items():
        sheet = wb.create_sheet(table)
        append_values(sheet, list(fields))
        for record in store.all(table):
            append_values(sheet, [record[key] for key in fields])
        style_sheet(sheet)
    save_workbook(wb, path)


def export_plan(plan, path):
    wb = Workbook()
    overview = wb.active
    overview.title = "Room Summary"
    append_values(overview, ["Room", "Supervisor ID", "Supervisor", "Exam date", "Start", "End",
                             "Students", "Seat capacity", "Benches used"])
    assignments = plan["assignments"]
    headers = ["Room", "Row", "Bench column", "Bench", "Side", "Roll number", "Name",
               "Year", "Division", "Subject code", "Subject"]
    all_seats = wb.create_sheet("All Students")
    append_values(all_seats, headers)
    student_rows = []
    for room in plan["rooms"]:
        seats = [a for a in assignments if a["room_no"] == room["room_no"]]
        total = sum(bool(a[side]) for a in seats for side in ("left", "right"))
        append_values(overview, [room["room_no"], room["supervisor"]["staff_id"],
            room["supervisor"]["name"], plan["date"], plan["start"], plan["end"], total,
            room["rows"] * room["benches_per_row"] * 2, len(seats)])
        sheet = wb.create_sheet(f"Room {room['room_no']}")
        append_values(sheet, ["Row", "Bench column", "Bench", "Left roll", "Left name", "Left subject",
                              "Right roll", "Right name", "Right subject"])
        by_bench = {a["bench"]: a for a in seats}
        for index in range(room["rows"] * room["benches_per_row"]):
            a = by_bench.get(index + 1)
            left, right = (a["left"], a["right"]) if a else (None, None)
            def detail(person):
                if person:
                    return [person["roll_no"], person["name"], person["subject_name"]]
                return ["EMPTY", "", ""]
            append_values(sheet, [index // room["benches_per_row"] + 1,
                index % room["benches_per_row"] + 1, index + 1] + detail(left) + detail(right))
            if a:
                for side in ("left", "right"):
                    person = a[side]
                    if person:
                        student_rows.append([room["room_no"], a["row"], a["column"], a["bench"],
                            side.title(), person["roll_no"], person["name"], YEARS[person["year"]],
                            person["division"], person["subject_code"], person["subject_name"]])
        style_sheet(sheet)
        sheet.oddHeader.center.text = f"Room {room['room_no']} | {plan['date']} {plan['start']}-{plan['end']}"
        sheet.oddFooter.center.text = f"Supervisor: {room['supervisor']['name']} | Page &P of &N"
    for row in sorted(student_rows, key=lambda r: r[5]):
        append_values(all_seats, row)
    style_sheet(overview)
    style_sheet(all_seats)
    save_workbook(wb, path)
